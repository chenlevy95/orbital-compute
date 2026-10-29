"""Load Finance's commitment tracker snapshots into SQLite with SCD Type 2 history.

Every commitments_YYYY-MM-DD.xlsx is parsed and validated before the database is
touched. New snapshots are stored as observations, then contract history is rebuilt
from all observations in snapshot-date order, so the arrival order of files does
not matter. An invalid file is skipped and reported, and the run exits 1; the other
files still load. A skipped file is never stored, so once it is fixed the next run
loads it and the rebuild puts it in place. A contract that disappears from a later
snapshot keeps its last known version.

Snapshot dates are observation dates, not amendment effective dates.
"""

import argparse
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
import sys

from openpyxl import load_workbook

import db

FIELDS = ('contract_id', 'customer_id', 'customer_name', 'start_date',
          'end_date', 'commitment_usd', 'status')

# Normalized workbook header -> field. Finance renames columns between exports;
# add an alias here once its meaning is confirmed. Unknown columns are ignored.
ALIASES = {
    'contract id': 'contract_id',
    'customer id': 'customer_id',
    'customer': 'customer_name',
    'start date': 'start_date',
    'end date': 'end_date',
    'total commitment (usd)': 'commitment_usd',
    'commit (usd)': 'commitment_usd',
    'status': 'status',
}


class TrackerError(Exception):
    """A source or business-rule problem, safe to print."""


def header(value):
    return ' '.join(str(value).split()).casefold() if value is not None else ''


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def parse_date(value):
    """Accept Excel dates, YYYY-MM-DD, and '1 Jun 2026'."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        for fmt in ('%Y-%m-%d', '%d %b %Y'):
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    raise ValueError(f'not a date: {value!r}')


def parse_money(value):
    """Return a positive USD amount as exact decimal text, e.g. '$900,000.00' -> '900000.00'."""
    if isinstance(value, str) and re.fullmatch(r'\s*\$?[\d,]+(\.\d+)?\s*', value):
        value = value.strip().lstrip('$').replace(',', '')
    if isinstance(value, bool) or not isinstance(value, (int, float, str)):
        raise ValueError(f'not an amount: {value!r}')
    try:
        amount = Decimal(str(value))
    except InvalidOperation:
        raise ValueError(f'not an amount: {value!r}') from None
    if not amount.is_finite() or amount <= 0:
        raise ValueError(f'amount must be positive: {value!r}')
    return str(amount.quantize(Decimal('0.01')))


def parse_text(value):
    if blank(value):
        raise ValueError('empty value')
    return str(value).strip()


PARSERS = {'start_date': parse_date, 'end_date': parse_date, 'commitment_usd': parse_money}


def snapshot_date(filename):
    match = re.fullmatch(r'commitments_(\d{4}-\d{2}-\d{2})\.xlsx', filename)
    if not match:
        raise TrackerError(f'{filename}: expected commitments_YYYY-MM-DD.xlsx')
    return date.fromisoformat(match[1]).isoformat()


def find_table(workbook, filename):
    """Return (sheet, header_row, {field: column index}) for the one row naming every field."""
    found = []
    for sheet in workbook:
        for row in sheet.iter_rows():
            columns = {}
            for index, cell in enumerate(row):
                field = ALIASES.get(header(cell.value))
                if field in columns:
                    raise TrackerError(f'{filename}:{sheet.title}:{cell.row}: two columns map to {field}')
                if field:
                    columns[field] = index
            if set(FIELDS) <= set(columns):
                found.append((sheet, row[0].row, columns))
    if len(found) != 1:
        raise TrackerError(f'{filename}: expected one header row with {", ".join(FIELDS)}; '
                           f'found {len(found)}. Check the layout and ALIASES.')
    return found[0]


def parse_snapshot(path):
    """Read one workbook and return (as_of, checksum, rows sorted by contract_id)."""
    path = Path(path)
    as_of = snapshot_date(path.name)
    content = path.read_bytes()
    workbook = load_workbook(BytesIO(content))
    try:
        sheet, header_row, columns = find_table(workbook, path.name)
        rows = {}
        for number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True),
                                       header_row + 1):
            values = {field: cells[index] for field, index in columns.items()}
            contract = values['contract_id']
            # Skip blank lines, the TOTAL row and '*' footnotes; none has a customer.
            if blank(values['customer_id']) and (
                    blank(contract) or str(contract).strip().upper() == 'TOTAL'
                    or str(contract).strip().startswith('*')):
                continue
            location = f'{path.name}:{sheet.title}:row {number}'
            try:
                row = tuple(PARSERS.get(field, parse_text)(values[field]) for field in FIELDS)
            except ValueError as exc:
                raise TrackerError(f'{location}: {exc}') from None
            if row[3] > row[4]:
                raise TrackerError(f'{location}: end date before start date')
            if rows.get(row[0], row) != row:
                raise TrackerError(f'{path.name}: conflicting rows for contract {row[0]}')
            rows[row[0]] = row  # an exact duplicate row collapses into one
    finally:
        workbook.close()
    if not rows:
        raise TrackerError(f'{path.name}: no contract rows')
    return as_of, sha256(content).hexdigest(), [rows[key] for key in sorted(rows)]


def create_schema(connection):
    columns = ', '.join(f'{field} TEXT NOT NULL' for field in FIELDS)
    connection.execute('''CREATE TABLE IF NOT EXISTS commitment_snapshots (
        as_of TEXT PRIMARY KEY, filename TEXT NOT NULL, checksum TEXT NOT NULL)''')
    connection.execute(f'''CREATE TABLE IF NOT EXISTS commitment_observations (
        as_of TEXT NOT NULL REFERENCES commitment_snapshots(as_of), {columns},
        PRIMARY KEY (as_of, contract_id))''')
    # [observed_from, observed_to); observed_to IS NULL marks the latest known version.
    connection.execute(f'''CREATE TABLE IF NOT EXISTS contract_history (
        {columns}, observed_from TEXT NOT NULL, observed_to TEXT,
        PRIMARY KEY (contract_id, observed_from))''')
    connection.execute('''CREATE VIEW IF NOT EXISTS contracts_current AS
        SELECT * FROM contract_history WHERE observed_to IS NULL''')


def check_no_overlap(current, as_of):
    """End dates are inclusive, so a customer's next contract must start after the previous end."""
    ordered = sorted(current.values(), key=lambda row: (row[1], row[3]))
    for previous, row in zip(ordered, ordered[1:]):
        if previous[1] == row[1] and row[3] <= previous[4]:
            raise TrackerError(f'as of {as_of}: {previous[0]} and {row[0]} overlap '
                               f'for customer {row[1]}')


def rebuild_history(connection):
    connection.execute('DELETE FROM contract_history')
    current = {}
    snapshots = [as_of for (as_of,) in connection.execute(
        'SELECT as_of FROM commitment_snapshots ORDER BY as_of')]
    for as_of in snapshots:
        for row in connection.execute(f'SELECT {", ".join(FIELDS)} FROM commitment_observations '
                                      'WHERE as_of=? ORDER BY contract_id', (as_of,)).fetchall():
            if current.get(row[0]) == row:
                continue
            connection.execute('UPDATE contract_history SET observed_to=? '
                               'WHERE contract_id=? AND observed_to IS NULL', (as_of, row[0]))
            connection.execute(f'INSERT INTO contract_history VALUES ({", ".join("?" * len(FIELDS))}, ?, NULL)',
                               (*row, as_of))
            current[row[0]] = row
        check_no_overlap(current, as_of)


def add_snapshot(connection, filename, as_of, checksum, rows):
    """Store a new snapshot's rows; return False if this exact file is already stored."""
    stored = connection.execute('SELECT checksum FROM commitment_snapshots WHERE as_of=?',
                                (as_of,)).fetchone()
    if stored:
        if stored[0] != checksum:
            raise TrackerError(f'{filename}: snapshot {as_of} was already loaded '
                               'with different contents')
        return False
    connection.execute('INSERT INTO commitment_snapshots VALUES (?, ?, ?)',
                       (as_of, filename, checksum))
    connection.executemany(
        f'INSERT INTO commitment_observations VALUES (?, {", ".join("?" * len(FIELDS))})',
        [(as_of, *row) for row in rows])
    return True


def ingest(input_dir, database):
    paths = sorted(Path(input_dir).glob('commitments_*.xlsx'))
    if not paths:
        raise TrackerError(f'{input_dir}: no commitments_YYYY-MM-DD.xlsx files')
    # Parse before taking the write lock. A file that fails is skipped, not stored.
    snapshots, errors = [], []
    for path in paths:
        try:
            snapshots.append((path.name, *parse_snapshot(path)))
        except TrackerError as exc:
            errors.append(str(exc))

    connection = db.connect(database)
    try:
        connection.execute('BEGIN IMMEDIATE')
        create_schema(connection)
        # a counter of new snapshots successfully added during this ingestion run
        added = 0
        for filename, as_of, checksum, rows in snapshots:
            # Some checks (overlaps, changed contents) need the stored data, so load
            # each file in a savepoint and undo only that file if one fails.
            connection.execute('SAVEPOINT snapshot')
            try:
                if add_snapshot(connection, filename, as_of, checksum, rows):
                    # An older snapshot arriving late can change later versions, so rebuild all.
                    rebuild_history(connection)
                    added += 1
            except TrackerError as exc:
                connection.execute('ROLLBACK TO snapshot')
                errors.append(str(exc))
            connection.execute('RELEASE snapshot')
        count = lambda table: connection.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0]
        result = {'snapshots_added': added, 'snapshots': count('commitment_snapshots'),
                  'history_versions': count('contract_history'),
                  'contracts': count('contracts_current'), 'errors': errors}
        connection.commit()
        return result
    except BaseException:
        connection.rollback()
        raise
    finally:
        connection.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    parser.add_argument('--input-dir', default='data/commitments')
    parser.add_argument('--database', default=os.getenv('USAGE_DATABASE', 'storage/usage.sqlite3'))
    args = parser.parse_args()
    try:
        result = ingest(args.input_dir, args.database)
        print(json.dumps(result, sort_keys=True))
        for error in result['errors']:
            print(f'Skipped: {error}', file=sys.stderr)
        return 1 if result['errors'] else 0
    except TrackerError as exc:
        print(f'Tracker ingestion failed: {exc}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
